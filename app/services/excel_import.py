"""Excel/CSV price list import & export.

Workflow in the admin dashboard:
  Download template (or "Export current prices") -> edit in Excel -> upload -> preview changes -> Apply.
Applying creates missing providers/services, updates prices (recording history and notifying
users), and the customer search results reflect the new prices immediately.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from ..models import ImportJob, Provider, ProviderService, Service, User, utcnow
from ..utils import PROVIDER_TYPES, SERVICE_CATEGORIES, parse_price, unique_slug
from .pricing import get_or_create_service, upsert_listing

COLUMNS = [
    ("provider_id", "Provider ID", "Optional. Leave blank for new providers."),
    ("provider_name", "Provider Name", "Required."),
    ("provider_type", "Provider Type", "e.g. Hospital, Clinic, Dental Practice."),
    ("address", "Address", ""),
    ("city", "City", "Used to match providers with the same name."),
    ("province", "Province", ""),
    ("phone", "Phone", ""),
    ("email", "Email", ""),
    ("website", "Website", ""),
    ("booking_url", "Booking URL", "Where the Book Now button sends patients."),
    ("google_profile_url", "Google Business Profile URL", "Paste the Share link from Google Maps. Reviews & location sync automatically."),
    ("avg_wait_minutes", "Avg Wait (minutes)", ""),
    ("service_name", "Service Name", "Required. New names are added to the catalogue."),
    ("service_category", "Service Category", "e.g. Dental, Imaging, Surgery."),
    ("price_min", "Price Min (ZAR)", "Required. Numbers only, e.g. 1250."),
    ("price_max", "Price Max (ZAR)", "Leave blank for a fixed price."),
    ("available", "Available", "Yes / No."),
    ("notes", "Notes", "Shown to patients, e.g. 'Excludes anaesthetist fee'."),
]
REQUIRED = ["provider_name", "service_name", "price_min"]
PROVIDER_FIELDS = ["provider_type", "address", "city", "province", "phone", "email", "website", "booking_url"]

HEADER_ALIASES = {
    "id": "provider_id", "provider id": "provider_id",
    "provider": "provider_name", "provider name": "provider_name", "facility": "provider_name",
    "facility name": "provider_name", "hospital": "provider_name", "practice": "provider_name", "name": "provider_name",
    "type": "provider_type", "provider type": "provider_type", "facility type": "provider_type",
    "service": "service_name", "service name": "service_name", "procedure": "service_name", "treatment": "service_name",
    "category": "service_category", "service category": "service_category",
    "price": "price", "price zar": "price", "cost": "price",
    "price min": "price_min", "min price": "price_min", "price from": "price_min", "from": "price_min", "minimum price": "price_min",
    "price max": "price_max", "max price": "price_max", "price to": "price_max", "to": "price_max", "maximum price": "price_max",
    "google business profile url": "google_profile_url", "google profile": "google_profile_url",
    "google business profile": "google_profile_url", "google maps link": "google_profile_url", "google url": "google_profile_url",
    "avg wait minutes": "avg_wait_minutes", "wait minutes": "avg_wait_minutes", "wait time": "avg_wait_minutes",
    "booking url": "booking_url", "booking link": "booking_url",
    "available": "available", "availability": "available", "notes": "notes", "comments": "notes",
    "address": "address", "city": "city", "town": "city", "province": "province", "phone": "phone",
    "telephone": "phone", "email": "email", "website": "website",
}


def _normalise_header(value) -> str | None:
    text = re.sub(r"\(.*?\)", " ", str(value or "")).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return HEADER_ALIASES.get(text) or (text.replace(" ", "_") if text.replace(" ", "_") in dict((c[0], 1) for c in COLUMNS) else None)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _truthy(value: str) -> bool | None:
    v = value.strip().lower()
    if not v:
        return None
    if v in {"yes", "y", "true", "1", "available"}:
        return True
    if v in {"no", "n", "false", "0", "unavailable"}:
        return False
    return None


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def read_table(content: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read the spreadsheet ({exc}). Upload an .xlsx or .csv file.") from exc
    sheet = wb["Prices"] if "Prices" in wb.sheetnames else wb.worksheets[0]
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def parse_rows(table: list[list]) -> ParseResult:
    result = ParseResult()
    header_idx = next(
        (i for i, row in enumerate(table[:10]) if sum(1 for v in row if _normalise_header(v)) >= 2), None
    )
    if header_idx is None:
        result.errors.append({"row": 1, "error": "Could not find a header row with Provider Name, Service Name and Price columns."})
        return result
    mapping = {i: _normalise_header(v) for i, v in enumerate(table[header_idx])}
    found = set(mapping.values())
    missing = [c for c in REQUIRED if c not in found and not (c == "price_min" and "price" in found)]
    if missing:
        labels = {k: label for k, label, _ in COLUMNS}
        result.errors.append({"row": header_idx + 1, "error": "Missing required column(s): " + ", ".join(labels[m] for m in missing)})
        return result

    for offset, raw in enumerate(table[header_idx + 1:], start=header_idx + 2):
        if not any(_cell(v) for v in raw):
            continue
        row = {key: _cell(raw[i]) for i, key in mapping.items() if key and i < len(raw)}
        if row.get("price") and not row.get("price_min"):
            row["price_min"] = row["price"]
        problems = []
        if not row.get("provider_name"):
            problems.append("Provider Name is empty")
        if not row.get("service_name"):
            problems.append("Service Name is empty")
        price_min = parse_price(row.get("price_min"))
        price_max = parse_price(row.get("price_max")) if row.get("price_max") else price_min
        if price_min is None:
            problems.append(f"Price Min '{row.get('price_min', '')}' is not a number")
        elif price_max is None:
            problems.append(f"Price Max '{row.get('price_max')}' is not a number")
        elif price_min < 0 or price_max < 0:
            problems.append("Prices cannot be negative")
        wait = row.get("avg_wait_minutes", "")
        if wait and not re.fullmatch(r"\d+(\.0+)?", wait):
            problems.append(f"Avg Wait '{wait}' must be a whole number of minutes")
        if problems:
            result.errors.append({"row": offset, "error": "; ".join(problems)})
            continue
        if price_min > price_max:
            price_min, price_max = price_max, price_min
        row.update(
            row=offset,
            price_min=price_min,
            price_max=price_max,
            available=_truthy(row.get("available", "")),
            avg_wait_minutes=int(float(wait)) if wait else None,
        )
        row.pop("price", None)
        result.rows.append(row)
    return result


def _find_provider(db: Session, row: dict) -> Provider | None:
    if row.get("provider_id", "").isdigit():
        provider = db.get(Provider, int(row["provider_id"]))
        if provider:
            return provider
    name = row["provider_name"].strip().lower()
    stmt = select(Provider).where(func.lower(Provider.name) == name)
    if row.get("city"):
        match = db.scalar(stmt.where(func.lower(Provider.city) == row["city"].strip().lower()))
        if match:
            return match
    matches = db.scalars(stmt).all()
    return matches[0] if len(matches) == 1 else None


def build_preview(db: Session, rows: list[dict]) -> dict:
    summary = {"rows": len(rows), "new_providers": 0, "new_services": 0, "new_prices": 0,
               "price_changes": 0, "unchanged": 0, "google_links": 0}
    seen_new_providers: set[tuple] = set()
    seen_new_services: set[str] = set()
    for row in rows:
        provider = _find_provider(db, row)
        key = (row["provider_name"].lower(), row.get("city", "").lower())
        if provider is None:
            row["provider_action"] = "new"
            if key not in seen_new_providers:
                seen_new_providers.add(key)
                summary["new_providers"] += 1
        else:
            row["provider_action"] = "existing"
            row["matched_provider_id"] = provider.id
            if row.get("google_profile_url") and row["google_profile_url"] != provider.google_profile_url:
                summary["google_links"] += 1
        service = db.scalar(select(Service).where(Service.name.ilike(" ".join(row["service_name"].split()))))
        if service is None:
            row["service_action"] = "new"
            if row["service_name"].lower() not in seen_new_services:
                seen_new_services.add(row["service_name"].lower())
                summary["new_services"] += 1
        else:
            row["service_action"] = "existing"
        listing = None
        if provider and service:
            listing = db.scalar(select(ProviderService).where(
                ProviderService.provider_id == provider.id, ProviderService.service_id == service.id))
        if listing is None:
            row["price_action"] = "create"
            summary["new_prices"] += 1
        elif abs(listing.price_min - row["price_min"]) > 0.005 or abs(listing.price_max - row["price_max"]) > 0.005:
            row["price_action"] = "update"
            row["old_price_min"], row["old_price_max"] = listing.price_min, listing.price_max
            summary["price_changes"] += 1
        else:
            row["price_action"] = "unchanged"
            summary["unchanged"] += 1
        if provider is None and row.get("google_profile_url"):
            summary["google_links"] += 1
    return summary


def create_import_job(db: Session, content: bytes, filename: str, user: User | None) -> ImportJob:
    parsed = parse_rows(read_table(content, filename))
    summary = build_preview(db, parsed.rows)
    summary["errors"] = len(parsed.errors)
    job = ImportJob(
        filename=filename,
        rows_json=json.dumps(parsed.rows),
        errors_json=json.dumps(parsed.errors),
        summary_json=json.dumps(summary),
        uploaded_by=user.id if user else None,
    )
    db.add(job)
    db.commit()
    return job


def apply_import_job(db: Session, job: ImportJob, *, hide_missing: bool = False) -> dict:
    """Apply a previewed job. Returns a result summary including provider ids needing a Google sync."""
    if job.status != "pending":
        raise ValueError(f"This import was already {job.status}.")
    rows = json.loads(job.rows_json)
    result = {"created": 0, "updated": 0, "unchanged": 0, "providers_created": 0, "services_created": 0,
              "hidden": 0, "google_sync_provider_ids": []}
    touched: dict[int, set[int]] = {}
    google_ids: set[int] = set()

    for row in rows:
        provider = _find_provider(db, row)
        if provider is None:
            provider = Provider(
                name=row["provider_name"].strip(),
                slug=unique_slug(db, Provider, f"{row['provider_name']} {row.get('city', '')}"),
            )
            db.add(provider)
            result["providers_created"] += 1
        for attr in PROVIDER_FIELDS:
            if row.get(attr):
                value = row[attr]
                if attr == "provider_type":
                    value = next((t for t in PROVIDER_TYPES if t.lower() == value.lower()), value)
                setattr(provider, attr, value)
        if row.get("avg_wait_minutes") is not None:
            provider.avg_wait_minutes = row["avg_wait_minutes"]
        link = row.get("google_profile_url", "")
        if link and link != provider.google_profile_url:
            provider.google_profile_url = link
            provider.google_place_id = ""  # re-resolve from the new link
            db.flush()
            google_ids.add(provider.id)
        db.flush()

        category = row.get("service_category") or None
        if category:
            category = next((c for c in SERVICE_CATEGORIES if c.lower() == category.lower()), category)
        service, created = get_or_create_service(db, row["service_name"], category)
        result["services_created"] += int(created)
        _, status = upsert_listing(
            db, provider, service, row["price_min"], row["price_max"], source="excel",
            notes=row.get("notes") if row.get("notes") else None, is_available=row.get("available"),
        )
        result[status] += 1
        touched.setdefault(provider.id, set()).add(service.id)

    if hide_missing:
        for provider_id, service_ids in touched.items():
            for listing in db.scalars(select(ProviderService).where(
                ProviderService.provider_id == provider_id,
                ProviderService.service_id.not_in(service_ids),
                ProviderService.is_available.is_(True),
            )):
                listing.is_available = False
                result["hidden"] += 1

    job.status = "applied"
    job.applied_at = utcnow()
    summary = json.loads(job.summary_json)
    summary["result"] = result
    job.summary_json = json.dumps(summary)
    db.commit()
    result["google_sync_provider_ids"] = sorted(google_ids)
    return result


# --------------------------------------------------------------------------- template / export

HEADER_FILL = PatternFill("solid", fgColor="0F766E")


def _styled_sheet(wb: Workbook) -> "Worksheet":  # noqa: F821
    ws = wb.active
    ws.title = "Prices"
    ws.append([label for _, label, _ in COLUMNS])
    for idx, (key, label, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        width = 42 if key in {"google_profile_url", "address", "notes"} else max(14, len(label) + 4)
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    avail_col = [i for i, c in enumerate(COLUMNS, start=1) if c[0] == "available"][0]
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(row=2, column=avail_col).column_letter}2:{ws.cell(row=2, column=avail_col).column_letter}5000")
    return ws


def _instructions(wb: Workbook, db: Session) -> None:
    ws = wb.create_sheet("Instructions")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    ws.append(["CostCare price list upload"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["One row = one service at one provider. Edit the 'Prices' sheet, save, and upload it in Admin → Import prices."])
    ws.append([])
    ws.append(["Column", "How to fill it"])
    ws["A4"].font = ws["B4"].font = Font(bold=True)
    for _, label, help_text in COLUMNS:
        ws.append([label, help_text])
    ws.append([])
    ws.append(["Provider types", ", ".join(PROVIDER_TYPES)])
    ws.append(["Service categories", ", ".join(SERVICE_CATEGORIES)])

    svc = wb.create_sheet("Service catalogue")
    svc.append(["Service Name", "Category", "Keywords"])
    for cell in svc[1]:
        cell.font = Font(bold=True)
    svc.column_dimensions["A"].width = 40
    svc.column_dimensions["B"].width = 18
    svc.column_dimensions["C"].width = 70
    for s in db.scalars(select(Service).order_by(Service.category, Service.name)):
        svc.append([s.name, s.category, s.keywords])


def build_template(db: Session) -> bytes:
    wb = Workbook()
    ws = _styled_sheet(wb)
    ws.append(["", "Example Day Clinic", "Clinic", "12 Long Street", "Cape Town", "Western Cape", "021 000 0000",
               "info@example.co.za", "https://example.co.za", "https://example.co.za/book",
               "https://maps.app.goo.gl/your-share-link", 20, "GP consultation", "Primary Care", 520, 780, "Yes",
               "Includes basic vitals check"])
    ws.append(["", "Example Day Clinic", "", "", "Cape Town", "", "", "", "", "", "", "", "Wound suturing",
               "Emergency", 950, 2200, "Yes", "Price depends on wound size"])
    _instructions(wb, db)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_prices(db: Session) -> bytes:
    wb = Workbook()
    ws = _styled_sheet(wb)
    listings = db.scalars(
        select(ProviderService)
        .join(ProviderService.provider)
        .join(ProviderService.service)
        .options(selectinload(ProviderService.provider), selectinload(ProviderService.service))
        .order_by(Provider.name, Service.name)
    )
    for l in listings:
        p = l.provider
        ws.append([p.id, p.name, p.provider_type, p.address, p.city, p.province, p.phone, p.email, p.website,
                   p.booking_url, p.google_profile_url, p.avg_wait_minutes, l.service.name, l.service.category,
                   l.price_min, l.price_max, "Yes" if l.is_available else "No", l.notes])
    _instructions(wb, db)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
