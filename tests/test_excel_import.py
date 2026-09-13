import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import Favorite, Notification, Provider, ProviderService, Service, User
from app.services import excel_import

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ["Provider Name", "City", "Provider Type", "Service Name", "Service Category",
           "Price Min (ZAR)", "Price Max (ZAR)", "Available", "Notes"]


def make_xlsx(rows, headers=HEADERS) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_rows_accepts_aliases_and_reports_errors():
    table = [
        ["Facility", "Procedure", "Price"],
        ["A clinic", "GP consultation", "R 1 250,00"],
        ["", "X-ray", "10"],
        ["B clinic", "Y", "abc"],
    ]
    result = excel_import.parse_rows(table)
    assert len(result.rows) == 1
    assert result.rows[0]["price_min"] == result.rows[0]["price_max"] == 1250.0
    assert [e["row"] for e in result.errors] == [3, 4]


def test_missing_required_column():
    result = excel_import.parse_rows([["Provider Name", "City"], ["A", "B"]])
    assert not result.rows
    assert "Missing required column" in result.errors[0]["error"]


def test_upload_preview_apply_updates_customer_search(admin_client, db):
    provider = db.scalar(select(Provider).where(Provider.name == "Gardens Sample Smile Clinic"))
    listing = db.scalar(select(ProviderService).join(Service).where(
        ProviderService.provider_id == provider.id, Service.name == "Root canal treatment"))
    old_min = listing.price_min

    traveller = User(email="saved-provider@test.local", name="T", password_hash="x")
    db.add(traveller)
    db.flush()
    db.add(Favorite(user_id=traveller.id, provider_id=provider.id))
    db.commit()

    content = make_xlsx([
        ["Gardens Sample Smile Clinic", "Cape Town", "", "Root canal treatment", "Dental", 3999, 5999, "Yes", "Updated via Excel"],
        ["Stellenbosch Test Clinic", "Stellenbosch", "Clinic", "Travel vaccine bundle", "Travel Health", 1500, "", "Yes", ""],
        ["Broken row", "Cape Town", "", "GP consultation", "", "not-a-price", "", "", ""],
    ])
    r = admin_client.post("/admin/import", files={"file": ("prices.xlsx", content, XLSX)}, follow_redirects=False)
    assert r.status_code == 303
    job_url = r.headers["location"]

    preview = admin_client.get(job_url)
    assert preview.status_code == 200
    assert "Travel vaccine bundle" in preview.text and "not a number" in preview.text

    # Nothing changes before apply
    db.expire_all()
    assert db.get(ProviderService, listing.id).price_min == old_min

    r = admin_client.post(job_url + "/apply", follow_redirects=False)
    assert r.status_code == 303

    db.expire_all()
    updated = db.get(ProviderService, listing.id)
    assert (updated.price_min, updated.price_max, updated.source) == (3999, 5999, "excel")
    assert updated.notes == "Updated via Excel"
    assert updated.history[0].old_min == old_min
    assert db.scalar(select(Notification).where(Notification.user_id == traveller.id)) is not None

    new_provider = db.scalar(select(Provider).where(Provider.name == "Stellenbosch Test Clinic"))
    assert new_provider and new_provider.city == "Stellenbosch" and new_provider.provider_type == "Clinic"

    data = admin_client.get("/api/search", params={"q": "root canal cape town"}).json()
    assert ("Gardens Sample Smile Clinic", 3999) in {(i["provider"]["name"], i["price_min"]) for i in data["results"]}

    # Applying twice is rejected
    admin_client.post(job_url + "/apply", follow_redirects=False)
    db.expire_all()
    assert db.get(ProviderService, listing.id).price_min == 3999


def test_template_and_export_downloads(admin_client):
    for url in ("/admin/import/template.xlsx", "/admin/export/prices.xlsx"):
        r = admin_client.get(url)
        assert r.status_code == 200
        assert r.headers["content-type"] == XLSX
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames[:2] == ["Prices", "Instructions"]
        assert wb["Prices"]["B1"].value == "Provider Name"


def test_exported_file_round_trips_without_changes(admin_client):
    exported = admin_client.get("/admin/export/prices.xlsx").content
    r = admin_client.post("/api/admin/import/excel", files={"file": ("export.xlsx", exported, XLSX)})
    assert r.status_code == 200
    summary = r.json()["summary"]
    assert summary["errors"] == 0
    assert summary["price_changes"] == 0 and summary["new_providers"] == 0
    assert summary["unchanged"] == summary["rows"]
